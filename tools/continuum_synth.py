"""Shared Continuum contract validator, VN working calendar, and deterministic artifact helpers."""

from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import openpyxl
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

VN_TZ_OFFSET = "+07:00"
VN_TZ = timezone(timedelta(hours=7))

VN_HOLIDAYS_2026: dict[date, str] = {
    date(2026, 1, 1): "New Year's Day (Tết Dương lịch)",
    date(2026, 2, 16): "Lunar New Year Holiday (Tết Nguyên Đán Bính Ngọ)",
    date(2026, 2, 17): "Lunar New Year Day 1 (Mùng 1 Tết Bính Ngọ)",
    date(2026, 2, 18): "Lunar New Year Day 2 (Mùng 2 Tết Bính Ngọ)",
    date(2026, 2, 19): "Lunar New Year Day 3 (Mùng 3 Tết Bính Ngọ)",
    date(2026, 2, 20): "Lunar New Year Holiday (Mùng 4 Tết Bính Ngọ)",
    date(2026, 4, 27): "Hung Kings' Commemoration Compensatory Day (Nghỉ bù Giỗ Tổ Hùng Vương)",
    date(2026, 4, 30): "Reunification Day (Ngày Giải phóng miền Nam 30/4)",
    date(2026, 5, 1): "International Labor Day (Ngày Quốc tế Lao động 1/5)",
    date(2026, 9, 1): "National Day Holiday Window (Nghỉ lễ Quốc khánh)",
    date(2026, 9, 2): "Vietnam National Day (Ngày Quốc khánh 2/9)",
}

# Verified against DATN_BE/src/services/iam, task/workspace, and capture
IAM_USER_STATUSES = {"ACTIVE", "SUSPENDED", "PENDING_INVITE"}
IAM_ORG_PLANS = {"FREE", "ENTERPRISE"}
IAM_ORG_MEM_STATUSES = {"PENDING_INVITE", "ACTIVE", "SUSPENDED", "REMOVED"}
IAM_PROJ_MEM_STATUSES = {"ACTIVE", "INACTIVE"}
IAM_PROJ_VISIBILITIES = {"PRIVATE", "PUBLIC"}
IAM_PROJ_STATUSES = {"ACTIVE", "ARCHIVED"}
IAM_ROLE_CODES = {"ADMIN", "TEAM_LEADER", "MEMBER"}

PROJECT_CODE_RE = re.compile(r"^[A-Z0-9]{2,10}$")
TASK_KEY_RE = re.compile(r"^[A-Z0-9]{2,10}-[1-9][0-9]*$")
HEX24_RE = re.compile(r"^[0-9a-f]{24}$")
SAFE_EMAIL_RE = re.compile(r"^[a-z0-9._%+-]+@[a-z0-9.-]+\.invalid$")

TASK_TYPES = {"TASK", "BUG", "STORY"}
TASK_STATUSES = {"TODO", "IN_PROGRESS", "IN_REVIEW", "BLOCKED", "DONE", "CANCELLED"}
TASK_PRIORITIES = {"LOW", "MEDIUM", "HIGH"}
SPRINT_STATUSES = {"PLANNED", "ACTIVE", "COMPLETED"}
TASK_EVENT_TYPES = {
    "TASK_CREATED",
    "TASK_STATUS_TRANSITIONED",
    "TASK_UPDATED",
    "TASK_COMMENT_ADDED",
    "TASK_WORKLOG_ADDED",
}

ALLOWED_TASK_TRANSITIONS: dict[str, set[str]] = {
    "TODO": {"IN_PROGRESS", "BLOCKED", "CANCELLED"},
    "IN_PROGRESS": {"TODO", "IN_REVIEW", "BLOCKED", "DONE", "CANCELLED"},
    "IN_REVIEW": {"IN_PROGRESS", "BLOCKED", "DONE", "CANCELLED"},
    "BLOCKED": {"TODO", "IN_PROGRESS", "IN_REVIEW", "CANCELLED"},
    "DONE": {"IN_PROGRESS"},
    "CANCELLED": {"TODO"},
}

CAPTURE_DOC_TYPES = {
    "MVP",
    "DATABASE_DESIGN",
    "RSR",
    "ARCHITECTURE",
    "API_SPECIFICATION",
    "RESEARCH",
    "REQUIREMENT",
    "DESIGN",
    "TECHNICAL_SPECIFICATION",
    "OTHER",
}
CAPTURE_DOC_STATUSES = {
    "DRAFT",
    "PENDING_REVIEW",
    "APPROVED",
    "CHANGES_REQUESTED",
    "REJECTED",
    "ARCHIVED",
}
CAPTURE_REPORT_SCOPES = {"MEMBER", "LEADER"}
CAPTURE_REPORT_FREQUENCIES = {
    "DAILY",
    "EVERY_2_DAYS",
    "EVERY_3_DAYS",
    "EVERY_4_DAYS",
    "EVERY_5_DAYS",
    "EVERY_6_DAYS",
    "WEEKLY",
}
CAPTURE_REPORT_STATUSES = {
    "DRAFT",
    "PENDING_REVIEW",
    "APPROVED",
    "CHANGES_REQUESTED",
    "REJECTED",
}
VALID_WORK_MODES = {
    "regular_hours",
    "evening_overtime",
    "weekend_incident_ot",
    "oncall_incident_response",
    "automated_bot",
}


def synth_oid(namespace: str, kind: str, source_id: str) -> str:
    """Deterministic 24-hex synthetic ObjectId candidate (never proof of live DB state)."""
    digest = hashlib.sha256(f"continuum-synth/v1/{namespace}/{kind}/{source_id}".encode("utf-8")).hexdigest()
    return digest[:24]


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def sha256_text(text: str) -> str:
    normalized = text.strip().replace("\r\n", "\n")
    return sha256_bytes(normalized.encode("utf-8"))


def is_working_day(d: date, holidays: set[date] | dict[date, str] = VN_HOLIDAYS_2026) -> bool:
    return d.weekday() < 5 and d not in holidays


def build_working_days(start: date, end: date, holidays: set[date] | dict[date, str] = VN_HOLIDAYS_2026) -> list[date]:
    days: list[date] = []
    cur = start
    while cur <= end:
        if is_working_day(cur, holidays):
            days.append(cur)
        cur += timedelta(days=1)
    return days


def snap_working_day(
    d: date,
    start: date,
    end: date,
    prefer_forward: bool = True,
    holidays: set[date] | dict[date, str] = VN_HOLIDAYS_2026,
) -> date:
    clamped = min(max(d, start), end)
    if is_working_day(clamped, holidays):
        return clamped
    step = 1 if prefer_forward else -1
    for direction in (step, -step):
        cur = clamped
        for _ in range(14):
            cur = cur + timedelta(days=direction)
            if start <= cur <= end and is_working_day(cur, holidays):
                return cur
    raise ValueError(f"Cannot snap {d} into working range [{start}, {end}]")


def add_working_days(
    d: date,
    delta_days: int,
    start: date,
    end: date,
    cap_at_end: bool = True,
    holidays: set[date] | dict[date, str] = VN_HOLIDAYS_2026,
) -> date:
    cur = snap_working_day(d, start, end, prefer_forward=True, holidays=holidays)
    rem = max(0, delta_days)
    while rem > 0:
        cur += timedelta(days=1)
        if cap_at_end and cur > end:
            return snap_working_day(end, start, end, prefer_forward=False, holidays=holidays)
        if is_working_day(cur, holidays):
            rem -= 1
    return cur


def make_vn_ts(d: date, hour: int, minute: int, second: int = 0) -> str:
    if not (0 <= hour <= 23 and 0 <= minute <= 59 and 0 <= second <= 59):
        raise ValueError(f"Invalid clock fields: {hour}:{minute}:{second}")
    return f"{d.isoformat()}T{hour:02d}:{minute:02d}:{second:02d}{VN_TZ_OFFSET}"


def make_regular_ts(d: date, seed_num: int, slot: int = 0) -> str:
    """Deterministic timestamp inside 08:00-11:55 or 13:00-16:55 +07:00 (never 12:00-12:59)."""
    hours = (8, 9, 10, 11, 13, 14, 15, 16)
    hr = hours[(seed_num + slot * 3) % len(hours)]
    mn = 4 + ((seed_num * 7 + slot * 11) % 52)
    sc = 3 + ((seed_num * 13 + slot * 17) % 54)
    return make_vn_ts(d, hr, mn, sc)


def make_ot_ts(d: date, seed_num: int, slot: int = 0) -> str:
    """Deterministic weekday evening overtime timestamp inside 17:05-21:55 +07:00."""
    hours = (17, 18, 19, 20, 21)
    hr = hours[(seed_num + slot) % len(hours)]
    mn = 5 + ((seed_num * 5 + slot * 13) % 50)
    sc = 4 + ((seed_num * 9 + slot * 19) % 52)
    return make_vn_ts(d, hr, mn, sc)


def parse_vn_ts(ts: str) -> datetime:
    if not isinstance(ts, str) or not ts.endswith(VN_TZ_OFFSET):
        raise ValueError(f"Timestamp must carry explicit +07:00 offset: {ts!r}")
    dt = datetime.fromisoformat(ts)
    if dt.utcoffset() != timedelta(hours=7):
        raise ValueError(f"Unexpected UTC offset on {ts}")
    return dt


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="\n") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def canonicalize_zip_file(path: Path, fixed_dt: tuple[int, int, int, int, int, int] = (2026, 10, 16, 0, 0, 0)) -> None:
    """Normalize ZIP entry mtimes and OOXML core.xml timestamps for byte-identical determinism."""
    iso_stamp = f"{fixed_dt[0]:04d}-{fixed_dt[1]:02d}-{fixed_dt[2]:02d}T00:00:00Z".encode("ascii")
    raw_bytes = path.read_bytes()
    in_buf = io.BytesIO(raw_bytes)
    out_buf = io.BytesIO()
    with zipfile.ZipFile(in_buf, "r") as zf_in, zipfile.ZipFile(out_buf, "w", compression=zipfile.ZIP_DEFLATED) as zf_out:
        for name in sorted(zf_in.namelist()):
            data = zf_in.read(name)
            if name.endswith("core.xml"):
                data = re.sub(
                    rb"(<dcterms:created[^>]*>).*?(</dcterms:created>)",
                    rb"\g<1>" + iso_stamp + rb"\g<2>",
                    data,
                )
                data = re.sub(
                    rb"(<dcterms:modified[^>]*>).*?(</dcterms:modified>)",
                    rb"\g<1>" + iso_stamp + rb"\g<2>",
                    data,
                )
            zinfo = zipfile.ZipInfo(name, date_time=fixed_dt)
            zinfo.compress_type = zipfile.ZIP_DEFLATED
            zf_out.writestr(zinfo, data)
    path.write_bytes(out_buf.getvalue())


def _shade_docx_cell(cell: Any, hex_color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    props.append(shd)


def _set_docx_cell(cell: Any, text: str, bold: bool = False, color: str = "1E293B") -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def render_formal_docx(
    target_path: Path,
    org_banner: str,
    doc_id: str,
    title: str,
    summary: str,
    metadata_pairs: list[tuple[str, str, str, str]],
    sections: list[tuple[str, list[str]]],
    reference_table: list[list[str]] | None = None,
    fixed_dt: tuple[int, int, int, int, int, int] = (2026, 10, 16, 0, 0, 0),
) -> str:
    """Create a deterministic, multi-section DOCX rendition with metadata header, bullets, and table."""
    target_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(30, 41, 59)
    normal.paragraph_format.space_after = Pt(5)

    for style_name, pt_size in (("Title", 20), ("Heading 1", 13)):
        st = doc.styles[style_name]
        st.font.name = "Aptos"
        st.font.size = Pt(pt_size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string("0F2942")

    hdr = sec.header.paragraphs[0]
    hdr.text = f"{org_banner}   |   SYNTHETIC CONTINUUM BENCHMARK ARTIFACT"
    hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in hdr.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 116, 139)

    ftr = sec.footer.paragraphs[0]
    ftr.text = f"100% Synthetic Project Document  •  {doc_id}"
    ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in ftr.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 116, 139)

    p_title = doc.add_paragraph(style="Title")
    p_title.paragraph_format.space_after = Pt(4)
    p_title.add_run(title)

    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_after = Pt(10)
    r_sum = p_sum.add_run(summary)
    r_sum.italic = True
    r_sum.font.color.rgb = RGBColor(71, 85, 105)

    meta_tbl = doc.add_table(rows=0, cols=4)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_vals in metadata_pairs:
        cells = meta_tbl.add_row().cells
        for idx, val in enumerate(row_vals):
            is_label = idx in (0, 2)
            _set_docx_cell(cells[idx], val, bold=is_label, color="FFFFFF" if is_label else "1E293B")
            _shade_docx_cell(cells[idx], "1E4E68" if is_label else "F1F5F9")

    doc.add_paragraph()
    for heading, paragraphs in sections:
        doc.add_heading(heading, level=1)
        for para_text in paragraphs:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            p.add_run(para_text)

    if reference_table and len(reference_table) >= 2:
        doc.add_heading("Bảng Tham Chiếu / Specification Matrix", level=1)
        tbl = doc.add_table(rows=1, cols=len(reference_table[0]))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, col_title in enumerate(reference_table[0]):
            _set_docx_cell(tbl.rows[0].cells[j], col_title, bold=True, color="FFFFFF")
            _shade_docx_cell(tbl.rows[0].cells[j], "0F2942")
        for r_idx, row_data in enumerate(reference_table[1:], start=1):
            cells = tbl.add_row().cells
            for j, cell_val in enumerate(row_data):
                _set_docx_cell(cells[j], cell_val)
                _shade_docx_cell(cells[j], "F8FAFC" if r_idx % 2 == 0 else "FFFFFF")

    doc.save(target_path)
    canonicalize_zip_file(target_path, fixed_dt=fixed_dt)
    return sha256_file(target_path)


def build_excel_catalog(
    target_path: Path,
    title: str,
    subtitle: str,
    kpis: list[tuple[str, str]],
    status_breakdown: list[tuple[str, int]],
    sheets: list[tuple[str, list[dict[str, Any]], list[str]]],
    fixed_dt: tuple[int, int, int, int, int, int] = (2026, 10, 16, 0, 0, 0),
) -> str:
    """Create a deterministic filterable XLSX workbook for dataset inspection."""
    target_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    navy = "0F2942"
    zebra = "F1F5F9"
    light = "F8FAFC"
    border_col = "CBD5E1"

    f_title = Font(name="Aptos", size=15, bold=True, color=navy)
    f_sub = Font(name="Aptos", size=9.5, italic=True, color="475569")
    f_hdr = Font(name="Aptos", size=9.5, bold=True, color="FFFFFF")
    f_cell = Font(name="Aptos", size=9, color="1E293B")
    f_kpi_v = Font(name="Aptos", size=16, bold=True, color=navy)
    f_kpi_l = Font(name="Aptos", size=8.5, bold=True, color="475569")

    fill_hdr = PatternFill(start_color=navy, end_color=navy, fill_type="solid")
    fill_zebra = PatternFill(start_color=zebra, end_color=zebra, fill_type="solid")
    fill_kpi = PatternFill(start_color=light, end_color=light, fill_type="solid")
    thin = Side(border_style="thin", color=border_col)
    bdr = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws_ov = wb.active
    ws_ov.title = "Overview"
    ws_ov.views.sheetView[0].showGridLines = True
    ws_ov.cell(2, 2, title).font = f_title
    ws_ov.cell(3, 2, subtitle).font = f_sub

    for idx, (lbl, val) in enumerate(kpis, start=2):
        cv = ws_ov.cell(5, idx, val)
        cl = ws_ov.cell(6, idx, lbl)
        cv.font = f_kpi_v
        cl.font = f_kpi_l
        cv.fill = fill_kpi
        cl.fill = fill_kpi
        cv.alignment = Alignment(horizontal="center", vertical="center")
        cl.alignment = Alignment(horizontal="center", vertical="center")
        cv.border = bdr
        cl.border = bdr

    ws_ov.cell(8, 2, "Task Status").font = f_hdr
    ws_ov.cell(8, 2).fill = fill_hdr
    ws_ov.cell(8, 3, "Count").font = f_hdr
    ws_ov.cell(8, 3).fill = fill_hdr
    for r_i, (st_name, cnt) in enumerate(status_breakdown, start=9):
        c1 = ws_ov.cell(r_i, 2, st_name)
        c2 = ws_ov.cell(r_i, 3, cnt)
        c1.font = f_cell
        c2.font = f_cell
        c1.border = bdr
        c2.border = bdr

    if status_breakdown:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Continuum Tasks by Status"
        data_ref = Reference(ws_ov, min_col=3, min_row=8, max_row=8 + len(status_breakdown))
        cats_ref = Reference(ws_ov, min_col=2, min_row=9, max_row=8 + len(status_breakdown))
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        chart.height = 7
        chart.width = 13
        ws_ov.add_chart(chart, "E8")

    all_ws = [ws_ov]
    for sh_name, rows, cols in sheets:
        ws = wb.create_sheet(title=sh_name[:31])
        all_ws.append(ws)
        ws.views.sheetView[0].showGridLines = True
        for c_i, col in enumerate(cols, start=1):
            hc = ws.cell(1, c_i, col)
            hc.font = f_hdr
            hc.fill = fill_hdr
        for r_i, rec in enumerate(rows, start=2):
            for c_i, col in enumerate(cols, start=1):
                val = rec.get(col)
                if isinstance(val, (list, dict)):
                    val = json.dumps(val, ensure_ascii=False)
                cc = ws.cell(r_i, c_i, "" if val is None else str(val))
                cc.font = f_cell
                if r_i % 2 == 1:
                    cc.fill = fill_zebra
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(2, len(rows) + 1)}"

    for ws in all_ws:
        for col_cells in ws.columns:
            letter = get_column_letter(col_cells[0].column)
            if ws == ws_ov:
                ws.column_dimensions[letter].width = 18
            else:
                max_len = max(len(str(c.value or "")) for c in col_cells[:80])
                ws.column_dimensions[letter].width = max(min(max_len + 3, 48), 12)

    wb.save(target_path)
    canonicalize_zip_file(target_path, fixed_dt=fixed_dt)
    return sha256_file(target_path)


def list_package_files(root: Path) -> list[Path]:
    """Return sorted physical files inside package root excluding bytecode/OS caches."""
    files: list[Path] = []
    for p in root.rglob("*"):
        if not p.is_file():
            continue
        if "__pycache__" in p.parts or p.name.endswith((".pyc", ".pyo")) or p.name in (".DS_Store", "Thumbs.db"):
            continue
        files.append(p)
    files.sort(key=lambda x: x.relative_to(root).as_posix())
    return files


def build_file_inventory(root: Path, exclude_names: set[str] | None = None) -> list[dict[str, Any]]:
    exclude = exclude_names or {"manifest.json"}
    inventory: list[dict[str, Any]] = []
    for p in list_package_files(root):
        rel = p.relative_to(root).as_posix()
        if rel in exclude:
            continue
        inventory.append(
            {
                "path": rel,
                "size_bytes": p.stat().st_size,
                "sha256": sha256_file(p),
            }
        )
    return inventory


def validate_continuum_organization_bundle(
    bundle: dict[str, Any],
    holidays: set[date] | dict[date, str] = VN_HOLIDAYS_2026,
) -> list[str]:
    """
    Comprehensive Continuum contract + temporal/ACL/Group-Note/evaluation validator.
    Returns a list of human-readable validation error strings (empty iff valid).
    """
    errors: list[str] = []
    as_of: str = bundle["as_of_date"]
    as_of_end_ts = f"{as_of}T23:59:59{VN_TZ_OFFSET}"
    primary_org_id: str = bundle["primary_org_id"]

    orgs = bundle["organizations"]
    users = bundle["users"]
    org_memberships = bundle["organization_memberships"]
    role_assignments = bundle["role_assignments"]
    projects = bundle["projects"]
    project_memberships = bundle["project_memberships"]
    sprints = bundle.get("sprints", [])
    tasks = bundle["tasks"]
    task_comments = bundle["task_comments"]
    task_worklogs = bundle["task_worklogs"]
    task_events = bundle["task_events"]
    task_notes = bundle["task_notes"]
    project_documents = bundle["project_documents"]
    document_versions = bundle["document_versions"]
    work_report_settings = bundle["work_report_settings"]
    work_reports = bundle["work_reports"]
    disruptions = bundle.get("disruption_events", [])
    oncall_schedules = bundle.get("oncall_schedules", [])
    eval_cases = bundle["evaluation_cases"]
    id_mappings = bundle.get("id_mappings", [])

    # Forbidden entities check (Prompt §4: NO Team, NO AI Chat history, NO 4th Notebook entity)
    for forbidden_key in ("teams", "team_memberships", "ai_chat_history", "notebooks"):
        if bundle.get(forbidden_key):
            errors.append(f"Forbidden scope entity present: {forbidden_key}")

    org_by_id = {o["org_id"]: o for o in orgs}
    if len(org_by_id) != len(orgs) or primary_org_id not in org_by_id:
        errors.append("Duplicate org_id or missing primary_org_id")
    for o in orgs:
        if o.get("plan") not in IAM_ORG_PLANS:
            errors.append(f"Invalid org plan on {o['org_id']}: {o.get('plan')}")
        if not HEX24_RE.match(str(o.get("continuum_org_id", ""))):
            errors.append(f"Invalid continuum_org_id on {o['org_id']}")

    user_by_id = {u["user_id"]: u for u in users}
    if len(user_by_id) != len(users):
        errors.append("Duplicate user_id in users")
    seen_emails: set[str] = set()
    for u in users:
        email = str(u.get("email", ""))
        if not SAFE_EMAIL_RE.match(email) or email in seen_emails:
            errors.append(f"Unsafe, non-lowercase, or duplicate email on {u['user_id']}: {email}")
        seen_emails.add(email)
        if u.get("status") not in IAM_USER_STATUSES:
            errors.append(f"Invalid IAM user status on {u['user_id']}: {u.get('status')}")
        if not HEX24_RE.match(str(u.get("continuum_user_id", ""))):
            errors.append(f"Invalid continuum_user_id on {u['user_id']}")

    # Organization memberships
    active_org_users: dict[str, set[str]] = {}
    all_org_users: dict[str, set[str]] = {}
    user_left_at: dict[str, str] = {}
    for om in org_memberships:
        oid, uid, st = om.get("organization_id"), om.get("user_id"), om.get("status")
        if oid not in org_by_id or uid not in user_by_id or st not in IAM_ORG_MEM_STATUSES:
            errors.append(f"Invalid organization_membership: {om.get('membership_id')}")
            continue
        all_org_users.setdefault(oid, set()).add(uid)
        if st == "ACTIVE" and user_by_id[uid]["status"] == "ACTIVE":
            active_org_users.setdefault(oid, set()).add(uid)
        if om.get("left_at"):
            user_left_at[uid] = om["left_at"]
    for u in users:
        if u.get("left_at"):
            user_left_at[u["user_id"]] = u["left_at"]

    # Projects
    project_by_id = {p["project_id"]: p for p in projects}
    if len(project_by_id) != len(projects):
        errors.append("Duplicate project_id in projects")
    seen_proj_codes: set[tuple[str, str]] = set()
    for p in projects:
        oid = p.get("organization_id")
        code = str(p.get("code", ""))
        if oid not in org_by_id:
            errors.append(f"Project {p['project_id']} references unknown org {oid}")
        if not PROJECT_CODE_RE.match(code) or (oid, code) in seen_proj_codes:
            errors.append(f"Invalid or duplicate project code on {p['project_id']}: {code}")
        seen_proj_codes.add((oid, code))
        if p.get("visibility") not in IAM_PROJ_VISIBILITIES or p.get("status") not in IAM_PROJ_STATUSES:
            errors.append(f"Invalid visibility/status on project {p['project_id']}")
        if p.get("created_by") not in all_org_users.get(oid, set()):
            errors.append(f"Project creator {p.get('created_by')} outside org {oid} on {p['project_id']}")

    # Project memberships
    active_proj_members: dict[str, set[str]] = {}
    historical_proj_members: dict[str, set[str]] = {}
    for pm in project_memberships:
        pid, oid, uid, st = pm.get("project_id"), pm.get("organization_id"), pm.get("user_id"), pm.get("status")
        proj = project_by_id.get(pid)
        if not proj or proj["organization_id"] != oid or uid not in user_by_id or st not in IAM_PROJ_MEM_STATUSES:
            errors.append(f"Invalid project_membership: {pm.get('membership_id')}")
            continue
        if uid not in all_org_users.get(oid, set()):
            errors.append(f"Cross-organization project_membership forbidden: user {uid} in project {pid} ({oid})")
        historical_proj_members.setdefault(pid, set()).add(uid)
        if st == "ACTIVE" and uid in active_org_users.get(oid, set()):
            active_proj_members.setdefault(pid, set()).add(uid)

    # Role assignments
    org_admins: dict[str, set[str]] = {}
    project_leaders: dict[str, set[str]] = {}
    for ra in role_assignments:
        oid, pid, uid, rcode = ra.get("organization_id"), ra.get("project_id"), ra.get("user_id"), ra.get("role_code")
        if oid not in org_by_id or uid not in all_org_users.get(oid, set()) or rcode not in IAM_ROLE_CODES:
            errors.append(f"Invalid role_assignment: {ra.get('assignment_id')}")
            continue
        if pid is None:
            if rcode == "TEAM_LEADER":
                errors.append(f"TEAM_LEADER must be project-scoped: {ra.get('assignment_id')}")
            if rcode == "ADMIN" and uid in active_org_users.get(oid, set()):
                org_admins.setdefault(oid, set()).add(uid)
        else:
            proj = project_by_id.get(pid)
            if not proj or proj["organization_id"] != oid or uid not in historical_proj_members.get(pid, set()):
                errors.append(f"Project role_assignment outside project membership: {ra.get('assignment_id')}")
                continue
            if rcode == "ADMIN":
                errors.append(f"ADMIN role must be org-scoped (project_id=null): {ra.get('assignment_id')}")
            if rcode == "TEAM_LEADER" and uid in active_proj_members.get(pid, set()):
                project_leaders.setdefault(pid, set()).add(uid)

    # Unavailability windows (sick leave, PTO, storm outage) & On-call windows
    unavailable_windows: dict[str, list[tuple[str, str, str]]] = {}
    for ev in disruptions:
        uid = ev.get("affected_user_id")
        if uid and ev.get("unavailable_from") and ev.get("unavailable_to"):
            unavailable_windows.setdefault(uid, []).append(
                (ev["unavailable_from"], ev["unavailable_to"], ev["disruption_id"])
            )

    oncall_windows: list[tuple[str, str, str]] = []
    for oc in oncall_schedules:
        for field_u in ("primary_user_id", "secondary_user_id"):
            uid = oc.get(field_u)
            if uid:
                oncall_windows.append((uid, oc["start_at"], oc["end_at"]))

    def check_actor_timestamp(
        actor_id: str | None,
        ts: str,
        work_mode: str = "regular_hours",
        is_ot: bool = False,
        actor_type: str = "human",
        context_label: str = "",
    ) -> None:
        try:
            dt = parse_vn_ts(ts)
        except ValueError as exc:
            errors.append(f"{context_label}: {exc}")
            return
        if ts > as_of_end_ts:
            errors.append(f"{context_label}: timestamp {ts} exceeds as_of {as_of_end_ts}")
        if work_mode not in VALID_WORK_MODES:
            errors.append(f"{context_label}: invalid work_mode {work_mode!r}")
        if actor_type == "system_bot":
            if work_mode != "automated_bot" or is_ot:
                errors.append(f"{context_label}: system_bot must use work_mode='automated_bot' and is_ot=False")
            return
        if not actor_id or actor_id not in user_by_id:
            errors.append(f"{context_label}: unknown human actor {actor_id!r}")
            return
        left_day = user_left_at.get(actor_id)
        if left_day and ts[:10] > left_day:
            errors.append(f"{context_label}: human {actor_id} acted at {ts} after offboarding {left_day}")
        for u_start, u_end, dis_id in unavailable_windows.get(actor_id, []):
            if u_start <= ts <= u_end:
                errors.append(f"{context_label}: human {actor_id} acted at {ts} during unavailable window {dis_id} ({u_start}..{u_end})")

        d = dt.date()
        hr = dt.hour
        if work_mode in ("weekend_incident_ot", "oncall_incident_response"):
            has_oncall = any(u == actor_id and s <= ts <= e for (u, s, e) in oncall_windows)
            if not has_oncall:
                errors.append(f"{context_label}: off-hours incident action by {actor_id} at {ts} lacks scheduled on-call coverage")
            return

        if not is_working_day(d, holidays):
            errors.append(f"{context_label}: routine human action by {actor_id} on weekend/holiday {d.isoformat()}")
            return
        if hr == 12:
            errors.append(f"{context_label}: human action by {actor_id} during mandatory lunch break 12:00-13:00 ({ts})")
        elif 8 <= hr <= 11 or 13 <= hr <= 16:
            if is_ot or work_mode != "regular_hours":
                errors.append(f"{context_label}: regular-hour timestamp {ts} marked as OT/{work_mode}")
        elif 17 <= hr <= 21 or (hr == 22 and dt.minute == 0 and dt.second == 0):
            if not is_ot or work_mode != "evening_overtime":
                errors.append(f"{context_label}: evening timestamp {ts} must set is_overtime=True and work_mode='evening_overtime'")
        else:
            errors.append(f"{context_label}: human action outside 08:00-22:00 without on-call exception ({ts})")

    # Sprints
    sprint_by_id = {s["sprint_id"]: s for s in sprints}
    if len(sprint_by_id) != len(sprints):
        errors.append("Duplicate sprint_id in sprints")
    active_sprint_count: dict[str, int] = {}
    for s in sprints:
        pid = s.get("project_id")
        proj = project_by_id.get(pid)
        if not proj or proj["organization_id"] != s.get("organization_id"):
            errors.append(f"Sprint {s['sprint_id']} has invalid project/org")
            continue
        if proj.get("methodology") in ("WATERFALL", "KANBAN"):
            errors.append(f"Project {pid} ({proj.get('methodology')}) must not define sprints ({s['sprint_id']})")
        if s.get("status") not in SPRINT_STATUSES:
            errors.append(f"Invalid sprint status on {s['sprint_id']}: {s.get('status')}")
        if s.get("start_date", "") > s.get("end_date", ""):
            errors.append(f"Inverted sprint dates on {s['sprint_id']}")
        if s.get("status") == "ACTIVE":
            active_sprint_count[pid] = active_sprint_count.get(pid, 0) + 1
    for pid, cnt in active_sprint_count.items():
        if cnt > 1:
            errors.append(f"Project {pid} has {cnt} ACTIVE sprints (Continuum allows at most 1)")

    # Tasks
    task_by_id = {t["task_id"]: t for t in tasks}
    if len(task_by_id) != len(tasks):
        errors.append("Duplicate task_id in tasks")
    seen_task_keys: set[tuple[str, str]] = set()
    for t in tasks:
        tid = t["task_id"]
        pid = t.get("project_id")
        oid = t.get("organization_id")
        proj = project_by_id.get(pid)
        if not proj or proj["organization_id"] != oid:
            errors.append(f"Task {tid} has invalid project/org")
            continue
        key = str(t.get("key", ""))
        if not TASK_KEY_RE.match(key) or not key.startswith(f"{proj['code']}-") or (oid, key) in seen_task_keys:
            errors.append(f"Invalid or duplicate Continuum task key on {tid}: {key}")
        seen_task_keys.add((oid, key))
        if t.get("type") not in TASK_TYPES or t.get("status") not in TASK_STATUSES or t.get("priority") not in TASK_PRIORITIES:
            errors.append(f"Invalid type/status/priority on task {tid}")
        if not (1 <= len(str(t.get("title", ""))) <= 255):
            errors.append(f"Invalid title length on task {tid}")
        assignee = t.get("assignee_id")
        if assignee is not None and assignee not in historical_proj_members.get(pid, set()):
            errors.append(f"Task {tid} assignee {assignee} is not in project {pid}")
        if t.get("status") in ("IN_PROGRESS", "IN_REVIEW") and (
            assignee is None or assignee not in active_proj_members.get(pid, set())
        ):
            errors.append(f"Active open task {tid} ({t.get('status')}) requires an ACTIVE project member assignee")
        creator = t.get("created_by")
        if creator not in historical_proj_members.get(pid, set()) and creator not in org_admins.get(oid, set()):
            errors.append(f"Task {tid} creator {creator} lacks project/org permission")
        sp_id = t.get("sprint_id")
        if proj.get("methodology") in ("WATERFALL", "KANBAN") and sp_id is not None:
            errors.append(f"Task {tid} in {proj.get('methodology')} project {pid} must have sprint_id=null")
        if sp_id is not None:
            sp = sprint_by_id.get(sp_id)
            if not sp or sp["project_id"] != pid:
                errors.append(f"Task {tid} references invalid sprint {sp_id}")
            elif sp["status"] == "COMPLETED" and t["status"] not in ("DONE", "CANCELLED"):
                errors.append(f"Unfinished task {tid} ({t['status']}) cannot remain in COMPLETED sprint {sp_id}")
        steps = t.get("steps", [])
        if len(steps) > 50 or any(not isinstance(s.get("done"), bool) or not s.get("text") for s in steps):
            errors.append(f"Invalid checklist steps on task {tid}")
        labels = t.get("labels", [])
        if len(labels) > 20 or any(not (1 <= len(str(lb)) <= 40) for lb in labels):
            errors.append(f"Invalid labels on task {tid}")
        if (t["status"] == "DONE") != (t.get("completed_at") is not None):
            errors.append(f"Task {tid} completed_at must be set iff status == DONE")
        if t["created_at"] > t["updated_at"]:
            errors.append(f"Task {tid} created_at > updated_at")
        check_actor_timestamp(
            creator,
            t["created_at"],
            t.get("created_work_mode", "regular_hours"),
            t.get("created_is_overtime", False),
            t.get("created_actor_type", "human"),
            f"task.created_at[{tid}]",
        )

    # Check 1-level subtask hierarchy
    for t in tasks:
        parent_id = t.get("parent_id")
        if parent_id is not None:
            parent = task_by_id.get(parent_id)
            if not parent or parent["project_id"] != t["project_id"]:
                errors.append(f"Subtask {t['task_id']} references missing/cross-project parent {parent_id}")
            elif parent.get("parent_id") is not None:
                errors.append(f"Subtask {t['task_id']} violates Continuum 1-level max subtask nesting")

    # Task comments, worklogs, and history events
    for c in task_comments:
        t = task_by_id.get(c.get("task_id"))
        if not t or c.get("project_id") != t["project_id"] or c.get("organization_id") != t["organization_id"]:
            errors.append(f"Invalid task_comment linkage: {c.get('comment_id')}")
            continue
        if not (1 <= len(str(c.get("body", ""))) <= 5000):
            errors.append(f"Invalid body length on comment {c.get('comment_id')}")
        if c.get("actor_type", "human") == "human":
            if c.get("author_id") not in historical_proj_members.get(t["project_id"], set()) and c.get(
                "author_id"
            ) not in org_admins.get(t["organization_id"], set()):
                errors.append(f"Comment {c['comment_id']} human author {c.get('author_id')} outside project")
        if c["created_at"] < t["created_at"]:
            errors.append(f"Comment {c['comment_id']} predates task {t['task_id']}")
        check_actor_timestamp(
            c.get("author_id"),
            c["created_at"],
            c.get("work_mode", "regular_hours"),
            c.get("is_overtime", False),
            c.get("actor_type", "human"),
            f"task_comment[{c['comment_id']}]",
        )

    for wl in task_worklogs:
        t = task_by_id.get(wl.get("task_id"))
        if not t or wl.get("project_id") != t["project_id"] or wl.get("organization_id") != t["organization_id"]:
            errors.append(f"Invalid task_worklog linkage: {wl.get('worklog_id')}")
            continue
        mins = wl.get("minutes", 0)
        if not isinstance(mins, int) or not (1 <= mins <= 1440):
            errors.append(f"Worklog {wl['worklog_id']} minutes {mins} outside 1..1440")
        if wl.get("author_id") not in historical_proj_members.get(t["project_id"], set()):
            errors.append(f"Worklog {wl['worklog_id']} author outside project")
        if wl["logged_at"] < t["created_at"]:
            errors.append(f"Worklog {wl['worklog_id']} predates task {t['task_id']}")
        check_actor_timestamp(
            wl.get("author_id"),
            wl["logged_at"],
            wl.get("work_mode", "regular_hours"),
            wl.get("is_overtime", False),
            "human",
            f"task_worklog[{wl['worklog_id']}]",
        )

    events_by_task: dict[str, list[dict[str, Any]]] = {}
    for ev in task_events:
        t = task_by_id.get(ev.get("task_id"))
        if not t or ev.get("project_id") != t["project_id"] or ev.get("type") not in TASK_EVENT_TYPES:
            errors.append(f"Invalid task_event: {ev.get('event_id')}")
            continue
        events_by_task.setdefault(t["task_id"], []).append(ev)
        check_actor_timestamp(
            ev.get("actor_id"),
            ev["occurred_at"],
            ev.get("work_mode", "regular_hours"),
            ev.get("is_overtime", False),
            ev.get("actor_type", "human"),
            f"task_event[{ev['event_id']}]",
        )
        if ev["type"] == "TASK_STATUS_TRANSITIONED":
            f_st, t_st = ev.get("from_status"), ev.get("to_status")
            if f_st not in ALLOWED_TASK_TRANSITIONS or t_st not in ALLOWED_TASK_TRANSITIONS.get(f_st, set()):
                errors.append(f"Disallowed status transition {f_st}->{t_st} on event {ev['event_id']}")
            needs_reason = (
                f_st == "BLOCKED"
                or t_st in ("BLOCKED", "CANCELLED")
                or f_st in ("DONE", "CANCELLED")
            )
            if needs_reason and len(str(ev.get("reason") or "").strip()) < 5:
                errors.append(f"Transition {f_st}->{t_st} on {ev['event_id']} requires reason >= 5 chars")

    for tid, ev_list in events_by_task.items():
        seq = [e["occurred_at"] for e in ev_list]
        if seq != sorted(seq):
            errors.append(f"Non-monotonic task_events on {tid}")
        status_evs = [e for e in ev_list if e["type"] == "TASK_STATUS_TRANSITIONED"]
        if status_evs and status_evs[-1]["to_status"] != task_by_id[tid]["status"]:
            errors.append(f"Final status transition {status_evs[-1]['to_status']} != task status {task_by_id[tid]['status']} on {tid}")

    # Group Note 1: Task Notes
    if not task_notes:
        errors.append("Group Note Form 1 (Task Note) is empty")
    note_by_id = {n["note_id"]: n for n in task_notes}
    if len(note_by_id) != len(task_notes):
        errors.append("Duplicate note_id in task_notes")
    for n in task_notes:
        t = task_by_id.get(n.get("task_id"))
        if not t or n.get("project_id") != t["project_id"] or n.get("organization_id") != t["organization_id"]:
            errors.append(f"Invalid Task Note linkage: {n.get('note_id')}")
            continue
        if n.get("type") != "TASK_NOTE" or n.get("status") not in ("CONFIRMED", "ARCHIVED"):
            errors.append(f"Invalid Task Note type/status on {n['note_id']}")
        if n.get("knowledge_eligible") is not False:
            errors.append(f"Task Note {n['note_id']} must NEVER have knowledge_eligible=True")
        if n.get("author_id") not in historical_proj_members.get(t["project_id"], set()):
            errors.append(f"Task Note {n['note_id']} author outside project")
        if n["created_at"] < t["created_at"]:
            errors.append(f"Task Note {n['note_id']} predates task {t['task_id']}")
        check_actor_timestamp(
            n.get("author_id"),
            n["created_at"],
            n.get("work_mode", "regular_hours"),
            n.get("is_overtime", False),
            "human",
            f"task_note[{n['note_id']}]",
        )

    # Group Note 2: Project Documents & Versions
    if not project_documents or not document_versions:
        errors.append("Group Note Form 2 (Project Document / Versions) is empty")
    doc_by_id = {d["document_id"]: d for d in project_documents}
    if len(doc_by_id) != len(project_documents):
        errors.append("Duplicate document_id in project_documents")
    versions_by_doc: dict[str, list[dict[str, Any]]] = {}
    for dv in document_versions:
        d = doc_by_id.get(dv.get("document_id"))
        if not d or dv.get("project_id") != d["project_id"] or dv.get("organization_id") != d["organization_id"]:
            errors.append(f"Invalid document_version linkage: {dv.get('version_id')}")
            continue
        if dv.get("status") not in CAPTURE_DOC_STATUSES:
            errors.append(f"Invalid document_version status on {dv['version_id']}")
        if dv.get("author_id") != d["created_by"]:
            errors.append(f"Document version {dv['version_id']} author {dv.get('author_id')} != document owner {d['created_by']}")
        check_actor_timestamp(
            dv.get("author_id"),
            dv["created_at"],
            dv.get("work_mode", "regular_hours"),
            dv.get("is_overtime", False),
            "human",
            f"doc_version.created[{dv['version_id']}]",
        )
        if dv["status"] in ("APPROVED", "CHANGES_REQUESTED", "REJECTED"):
            rev_id = dv.get("reviewed_by")
            pid, oid = d["project_id"], d["organization_id"]
            can_rev = rev_id in project_leaders.get(pid, set()) or rev_id in org_admins.get(oid, set())
            if not rev_id or rev_id == dv["author_id"] or not can_rev:
                errors.append(f"Invalid or self document review on {dv['version_id']} by {rev_id}")
            if not dv.get("reviewed_at") or dv["reviewed_at"] < dv["created_at"]:
                errors.append(f"Invalid reviewed_at chronology on {dv['version_id']}")
            else:
                check_actor_timestamp(
                    rev_id,
                    dv["reviewed_at"],
                    dv.get("review_work_mode", "regular_hours"),
                    dv.get("review_is_overtime", False),
                    "human",
                    f"doc_version.reviewed[{dv['version_id']}]",
                )
        versions_by_doc.setdefault(d["document_id"], []).append(dv)

    for did, d in doc_by_id.items():
        pid, oid = d.get("project_id"), d.get("organization_id")
        if pid not in project_by_id or project_by_id[pid]["organization_id"] != oid:
            errors.append(f"Project Document {did} has invalid project/org")
            continue
        if d.get("document_type") not in CAPTURE_DOC_TYPES or d.get("status") not in CAPTURE_DOC_STATUSES:
            errors.append(f"Invalid document_type/status on {did}")
        if d.get("created_by") not in historical_proj_members.get(pid, set()) and d.get("created_by") not in org_admins.get(oid, set()):
            errors.append(f"Document {did} creator outside project/org")
        dvs = sorted(versions_by_doc.get(did, []), key=lambda x: x["version_number"])
        if not dvs:
            errors.append(f"Document {did} has zero document_versions")
            continue
        approved_nums = [v["version_number"] for v in dvs if v["status"] == "APPROVED"]
        expected_latest_app = max(approved_nums) if approved_nums else None
        if d.get("latest_approved_version_number") != expected_latest_app:
            errors.append(f"Document {did} latest_approved_version_number mismatch")
        expected_ke = (expected_latest_app is not None) and (d["status"] != "ARCHIVED")
        if d.get("knowledge_eligible") != expected_ke:
            errors.append(f"Document {did} knowledge_eligible mismatch (expected {expected_ke})")

    # Group Note 3: Work Report Settings & Work Reports
    if not work_reports or not work_report_settings:
        errors.append("Group Note Form 3 (Work Report / Settings) is empty")
    for ws in work_report_settings:
        pid, oid = ws.get("project_id"), ws.get("organization_id")
        if pid not in project_by_id or project_by_id[pid]["organization_id"] != oid:
            errors.append(f"Invalid work_report_setting project/org: {ws.get('setting_id')}")
        if ws.get("report_scope") not in CAPTURE_REPORT_SCOPES or ws.get("frequency") not in CAPTURE_REPORT_FREQUENCIES:
            errors.append(f"Invalid scope/frequency on work_report_setting {ws.get('setting_id')}")

    report_by_id = {r["report_id"]: r for r in work_reports}
    if len(report_by_id) != len(work_reports):
        errors.append("Duplicate report_id in work_reports")
    scopes_seen: set[str] = set()
    freqs_seen: set[str] = set()
    for r in work_reports:
        rid = r["report_id"]
        pid, oid = r.get("project_id"), r.get("organization_id")
        if pid not in project_by_id or project_by_id[pid]["organization_id"] != oid:
            errors.append(f"Report {rid} has invalid project/org")
            continue
        scope = r.get("report_scope")
        freq = r.get("frequency")
        st = r.get("status")
        author = r.get("member_id")
        rev = r.get("reviewed_by")
        if scope not in CAPTURE_REPORT_SCOPES or freq not in CAPTURE_REPORT_FREQUENCIES or st not in CAPTURE_REPORT_STATUSES:
            errors.append(f"Invalid scope/frequency/status on report {rid}")
            continue
        scopes_seen.add(scope)
        freqs_seen.add(freq)
        if r.get("period_start", "") > r.get("period_end", ""):
            errors.append(f"Inverted period on report {rid}")
        if r.get("knowledge_eligible") != (st == "APPROVED"):
            errors.append(f"Report {rid} knowledge_eligible must equal (status == 'APPROVED')")

        # Continuum Capture strict authorization check:
        is_author_leader = author in project_leaders.get(pid, set())
        is_author_manager = author in org_admins.get(oid, set())
        if scope == "MEMBER":
            if author not in historical_proj_members.get(pid, set()) or is_author_leader or is_author_manager:
                errors.append(f"MEMBER report {rid} must be authored by regular project MEMBER, got {author}")
        else:  # LEADER
            if not is_author_leader:
                errors.append(f"LEADER report {rid} must be authored by project TEAM_LEADER, got {author}")

        check_actor_timestamp(
            author,
            r["created_at"],
            r.get("work_mode", "regular_hours"),
            r.get("is_overtime", False),
            "human",
            f"work_report.created[{rid}]",
        )
        if st in ("APPROVED", "CHANGES_REQUESTED", "REJECTED"):
            if not rev or rev == author:
                errors.append(f"Reviewed report {rid} missing reviewer or self-reviewed")
            elif scope == "MEMBER":
                if rev not in project_leaders.get(pid, set()) and rev not in org_admins.get(oid, set()):
                    errors.append(f"MEMBER report {rid} reviewer {rev} lacks TEAM_LEADER/ADMIN permission")
            else:  # LEADER report MUST be reviewed strictly by Organization ADMIN Manager!
                if rev not in org_admins.get(oid, set()):
                    errors.append(f"LEADER report {rid} must be reviewed by Organization ADMIN Manager, got {rev}")
            if not r.get("submitted_at") or not r.get("reviewed_at") or not (r["created_at"] <= r["submitted_at"] <= r["reviewed_at"]):
                errors.append(f"Invalid chronology on reviewed report {rid}")
            else:
                check_actor_timestamp(
                    rev,
                    r["reviewed_at"],
                    r.get("review_work_mode", "regular_hours"),
                    r.get("review_is_overtime", False),
                    "human",
                    f"work_report.reviewed[{rid}]",
                )

        if len(r.get("document_ids", [])) > 20:
            errors.append(f"Report {rid} exceeds 20 document_ids")
        for did in r.get("document_ids", []):
            if did not in doc_by_id or doc_by_id[did]["project_id"] != pid:
                errors.append(f"Report {rid} references missing/cross-project document {did}")
        for tid in r.get("source_task_ids", []):
            if tid not in task_by_id or task_by_id[tid]["project_id"] != pid:
                errors.append(f"Report {rid} references missing/cross-project task {tid}")

    if scopes_seen != {"MEMBER", "LEADER"}:
        errors.append(f"Organization must cover both MEMBER and LEADER report flows, got {scopes_seen}")
    if not {"DAILY", "WEEKLY"}.issubset(freqs_seen):
        errors.append(f"Organization must cover both DAILY and WEEKLY report frequencies, got {freqs_seen}")

    # Evaluation cases & Dual-reviewer governance
    if not eval_cases:
        errors.append("Evaluation suite is empty")
    all_citeable_ids = (
        set(task_by_id)
        | set(note_by_id)
        | set(doc_by_id)
        | {f"{dv['document_id']}@v{dv['version_number']}" for dv in document_versions}
        | set(report_by_id)
        | {ev["disruption_id"] for ev in disruptions}
    )
    seen_case_ids: set[str] = set()
    has_deny_case = False
    for c in eval_cases:
        cid = c.get("case_id", "")
        if not cid or cid in seen_case_ids:
            errors.append(f"Missing or duplicate evaluation case_id: {cid}")
        seen_case_ids.add(cid)
        if c.get("review_status") != "unreviewed" or c.get("review_status_label_vi") != "chưa duyệt":
            errors.append(f"Unreviewed synthetic benchmark case {cid} must explicitly declare review_status='unreviewed' ('chưa duyệt')")
        if "reviewer_1" not in c or "reviewer_2" not in c or "adjudication_status" not in c:
            errors.append(f"Evaluation case {cid} missing dual-reviewer governance fields")
        if c.get("split") not in ("dev", "test"):
            errors.append(f"Evaluation case {cid} must declare split in ('dev', 'test')")
        s_ids = set(c.get("source_ids", []))
        forb_ids = set(c.get("must_not_cite", []))
        if s_ids & forb_ids:
            errors.append(f"Evaluation case {cid} marks source both allowed and forbidden")
        if any(sid not in all_citeable_ids for sid in s_ids):
            errors.append(f"Evaluation case {cid} cites unknown source in {s_ids - all_citeable_ids}")
        if c.get("case_category") == "access_denied" or str(c.get("expected_answer", "")).startswith("DENY"):
            has_deny_case = True
            if s_ids:
                errors.append(f"Access-denied evaluation case {cid} must have empty source_ids")
    if not has_deny_case:
        errors.append("Evaluation suite must include at least one access-denied (DENY) case")

    # Source -> Target ID mapping completeness
    if id_mappings:
        seen_src = set()
        for m in id_mappings:
            key = (m.get("entity_type"), m.get("source_id"))
            if key in seen_src:
                errors.append(f"Duplicate id_mapping entry for {key}")
            seen_src.add(key)
            if not HEX24_RE.match(str(m.get("target_continuum_id", ""))):
                errors.append(f"Invalid 24-hex target_continuum_id in id_mapping for {key}")

    return errors
